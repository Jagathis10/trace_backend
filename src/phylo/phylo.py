#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
import argparse
import subprocess
from Bio import AlignIO, SeqIO
import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill, Alignment
from openpyxl import Workbook


def run_mafft(input_fasta, cpu=8):
    out_msa = f"{os.path.splitext(input_fasta)[0]}.msa"
    cmd = [
        "mafft",
        "--thread", str(cpu),
        input_fasta,
    ]
    print(f" Running MAFFT:\n  {' '.join(cmd)}")
    with open(out_msa, "w") as out_handle:
        result = subprocess.run(
            cmd,
            stdout=out_handle,
            stderr=subprocess.PIPE,
            text=True,
        )
    if result.returncode != 0:
        print(" MAFFT error:")
        print(result.stderr)
        raise RuntimeError("MAFFT failed")
    print(f" MAFFT alignment written to {out_msa}")
    return out_msa


def make_phylip(msa_file):
    basename = os.path.splitext(msa_file)[0]
    phylip_file = f"{basename}.phy"

    print(f" Converting {msa_file} to {phylip_file} (PHYLIP-relaxed)")

    seen_ids = set()
    unique_records = []

    for i, record in enumerate(SeqIO.parse(msa_file, "fasta")):
        original_id = record.id
        if record.id in seen_ids:
            record.id = f"{original_id}_{i}"  # Make name unique
        seen_ids.add(record.id)
        record.description = ""  # PHYLIP relaxed doesn't need description
        unique_records.append(record)

    SeqIO.write(unique_records, phylip_file, "phylip-relaxed")
    print(f"PHYLIP alignment written to {phylip_file}")
    return phylip_file


def calculate_identity_pair(seq1, seq2):
    matches = 0
    valid = 0
    for a, b in zip(seq1, seq2):
        if a == "-" or b == "-":
            continue
        valid += 1
        if a == b:
            matches += 1
    if valid == 0:
        return 0.0
    return round((matches / valid) * 100.0, 2)


def calculate_identity_matrix(alignment):
    num_seqs = len(alignment)
    identity_matrix = np.zeros((num_seqs, num_seqs), dtype=float)

    print(f" Calculating identity matrix for {num_seqs} sequences")

    for i in range(num_seqs):
        seq_i = str(alignment[i].seq)
        identity_matrix[i, i] = 100.0
        for j in range(i + 1, num_seqs):
            seq_j = str(alignment[j].seq)
            pid = calculate_identity_pair(seq_i, seq_j)
            identity_matrix[i, j] = pid
            identity_matrix[j, i] = pid

    return identity_matrix


def process_dataframe(df: pd.DataFrame) -> Workbook:
    fills = {
        (99, 100): PatternFill("solid", fgColor="FFCCCC"),  # 99–100
        (97, 99): PatternFill("solid", fgColor="FFDAB3"),   # 97–99
        (95, 97): PatternFill("solid", fgColor="FFFF99"),   # 95–97
        (93, 95): PatternFill("solid", fgColor="B3FFB3"),   # 93–95
        (90, 93): PatternFill("solid", fgColor="B3CCFF"),   # 90–93
        (-1, 90): PatternFill("solid", fgColor="FFFFFF"),   # < 90
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "identity_matrix"
    ws.append([""] + list(df.columns))

    for idx, row in df.iterrows():
        ws.append([idx] + row.tolist())

    # Color cells
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if isinstance(cell.value, (int, float, float.__mro__[0])):
                for (low, high), fill in fills.items():
                    if (low <= cell.value < high) or (low == 99 and cell.value == 100):
                        cell.fill = fill
                        break

    # Rotate header text
    for cell in ws[1]:
        cell.alignment = Alignment(textRotation=90)

    # Upper triangle
    mask = np.triu(np.ones(df.shape), k=1)
    upper_df = df.where(mask == 1)
    ws2 = wb.create_sheet("identity_matrix-1")
    ws2.append([""] + list(upper_df.columns))

    for cell in ws2[1]:
        cell.alignment = Alignment(textRotation=90)

    for idx, row in upper_df.iterrows():
        ws2.append([idx] + row.tolist())

    # Center everything
    for sheet in [ws, ws2]:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    return wb


def run_iqtree2(phylip_file, cpu=8):
    prefix = os.path.splitext(phylip_file)[0]
    cmd = [
        "iqtree2",
        "-s", phylip_file,
        "-m", "GTR+F+R10",
        "--alrt", "1000",
        "-B", "1000",
        "-T", str(cpu),
        "--prefix", prefix,
    ]
    print(f" Running IQ-TREE2:\n  {' '.join(cmd)}")
    result = subprocess.run(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if result.returncode != 0:
        print(" IQ-TREE2 error:")
        print(result.stdout)
        print(result.stderr)
        raise RuntimeError("IQ-TREE2 failed")
    print(f" IQ-TREE2 finished for {phylip_file}")


def parse_arguments():
    parser = argparse.ArgumentParser(
        description="Phylogenetic analysis pipeline with MAFFT and IQ-TREE2"
    )
    # Short options -i and -t match your phylo.js spawn call
    parser.add_argument("-i", "--input_fasta", required=True, help="Input FASTA file path")
    parser.add_argument("-t", "--cpu", type=int, default=8, help="Number of CPUs to use")
    return parser.parse_args()


def main():
    args = parse_arguments()
    input_fasta = os.path.abspath(args.input_fasta)
    cpu = args.cpu

  
    print(" Phylo.py started")
    print(f"Input FASTA : {input_fasta}")
    print(f"CPUs        : {cpu}")
  

    
    msa = run_mafft(input_fasta, cpu=cpu)
    phylip = make_phylip(msa)
    #Identity matrix
    alignment = AlignIO.read(msa, "fasta")
    headers = [record.id for record in alignment]
    matrix = calculate_identity_matrix(alignment)
    df = pd.DataFrame(matrix, index=headers, columns=headers)
    wb = process_dataframe(df)

    # IMPORTANT: match your phylo.js expectation: <base>_identity.xlsx
    excel_path = f"{os.path.splitext(input_fasta)[0]}_identity.xlsx"
    print(f" Saving colored identity matrix to {excel_path}")
    wb.save(excel_path)

    #IQ-TREE
    run_iqtree2(phylip, cpu=cpu)

    print(" All steps completed.")
    

if __name__ == "__main__":
    main()
